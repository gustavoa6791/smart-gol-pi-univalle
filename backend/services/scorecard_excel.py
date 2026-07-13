"""Genera la tarjeta de árbitro en Excel (.xlsx).

Misma estructura que la tarjeta imprimible en HTML: por equipo, una tabla con
# (camiseta), Nombre, Goles (10 casillas + T.G.), A (2 casillas + T.A.) y R (1
casilla + T.R.). Con blank=True se omite el # y el nombre pre-cargados.
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

_HEADER_FILL = PatternFill(start_color="1E7A34", end_color="1E7A34", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(bold=True, size=14)
_THIN = Side(style="thin", color="000000")
_THICK = Side(style="medium", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_TOTAL_BORDER = Border(left=_THICK, right=_THICK, top=_THICK, bottom=_THICK)
_FOLD_SIDE = Side(style="mediumDashed", color="000000")
_FOLD_BORDER = Border(left=_THIN, right=_FOLD_SIDE, top=_THIN, bottom=_THIN)
_TOTAL_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
_CENTER = Alignment(horizontal="center", vertical="center")

GOAL_TALLY = 10
YELLOW_TALLY = 2
RED_TALLY = 1


def _col_letter(col: int) -> str:
    letters = ""
    while col > 0:
        col, rem = divmod(col - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _write_team_table(ws: Worksheet, start_row: int, team_name: str, players: list[dict], blank: bool) -> int:
    equipo_row = start_row
    header_row = start_row + 1

    ws.cell(row=equipo_row, column=1, value="Equipo:").font = _HEADER_FONT
    ws.cell(row=equipo_row, column=1).fill = _HEADER_FILL
    ws.merge_cells(start_row=equipo_row, start_column=1, end_row=equipo_row, end_column=2)
    ws.cell(row=equipo_row, column=3, value=team_name).font = _TITLE_FONT
    last_col_hint = 3 + GOAL_TALLY + 1 + YELLOW_TALLY + 1 + RED_TALLY + 1 - 1
    ws.merge_cells(start_row=equipo_row, start_column=3, end_row=equipo_row, end_column=last_col_hint)

    col = 3
    blocks = [("Goles", "T.G.", GOAL_TALLY), ("A", "T.A.", YELLOW_TALLY), ("R", "T.R.", RED_TALLY)]
    tally_cols: list[int] = []
    total_cols: list[int] = []

    for label, total_label, count in blocks:
        block_start = col
        for _ in range(count):
            tally_cols.append(col)
            col += 1
        if count > 1:
            ws.merge_cells(start_row=header_row, start_column=block_start, end_row=header_row, end_column=col - 1)
        cell = ws.cell(row=header_row, column=block_start, value=label)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER

        total_cell = ws.cell(row=header_row, column=col, value=total_label)
        total_cell.fill = _HEADER_FILL
        total_cell.font = _HEADER_FONT
        total_cell.alignment = _CENTER
        total_cell.border = _TOTAL_BORDER
        total_cols.append(col)
        col += 1

    for c, label in ((1, "#"), (2, "Nombre")):
        cell = ws.cell(row=header_row, column=c, value=label)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER
        cell.border = _FOLD_BORDER if c == 2 else _BORDER

    last_col = col - 1
    row = header_row + 1
    for player in players:
        ws.cell(row=row, column=1, value=None if blank else player["jersey_number"]).border = _BORDER
        ws.cell(row=row, column=1).alignment = _CENTER
        ws.cell(row=row, column=2, value=None if blank else player["name"]).border = _FOLD_BORDER
        for c in tally_cols:
            cell = ws.cell(row=row, column=c, value=None)
            cell.border = _BORDER
            cell.alignment = _CENTER
        for c in total_cols:
            cell = ws.cell(row=row, column=c, value=None)
            cell.border = _TOTAL_BORDER
            cell.fill = _TOTAL_FILL
            cell.alignment = _CENTER
        row += 1

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 26
    for c in range(3, last_col + 1):
        letter = _col_letter(c)
        ws.column_dimensions[letter].width = 9 if c in total_cols else 4

    return row + 2


def build_scorecard_workbook(
    round_label: str,
    home_team_name: str,
    home_players: list[dict],
    away_team_name: str,
    away_players: list[dict],
    blank: bool,
) -> bytes:
    """players: lista de {"jersey_number": int | None, "name": str}."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Tarjeta"

    ws.cell(row=1, column=1, value=f"Tarjeta de árbitro — {round_label}").font = Font(bold=True, size=16)

    next_row = _write_team_table(ws, 3, home_team_name, home_players, blank)
    _write_team_table(ws, next_row, away_team_name, away_players, blank)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
